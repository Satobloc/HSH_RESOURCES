```python
#!/usr/bin/env python3

"""
Unicode String-to-Glyph Assignment Tool

Spreadsheet format:

    GLYPH | NO. | STRING | DEF

Input spreadsheets may contain only:

    STRING | DEF

Usage:

    python unicode_assigner.py --import input.xlsx

Then interactively:

    > hello
    > world
    > A

Commands:

    :import filename.xlsx
    :save
    :count
    :show STRING
    :quit

Rules:

1. Multi-character strings receive the first unused eligible
   Unicode character.

2. A one-code-point string claims that exact Unicode character.

3. If that character is already assigned, the old assignment
   is moved to a fresh eligible Unicode character.

4. The old assignment keeps its STRING and DEF.

5. Every change is saved immediately.

Unicode data:
    Unicode 18.0.0
"""

import argparse
import os
import sys
import shutil
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

UNICODE_VERSION = "18.0.0"

UNICODE_BASE = (
    f"https://www.unicode.org/Public/{UNICODE_VERSION}/ucd"
)

UNICODE_DATA_URL = (
    f"{UNICODE_BASE}/UnicodeData.txt"
)

EAST_ASIAN_WIDTH_URL = (
    f"{UNICODE_BASE}/EastAsianWidth.txt"
)

DATA_DIRECTORY = Path(".unicode_data")

DEFAULT_WORKBOOK = Path("unicode_assignments.xlsx")

HEADERS = [
    "GLYPH",
    "NO.",
    "STRING",
    "DEF",
]


# ============================================================
# GENERAL HELPERS
# ============================================================

def download_file(url, destination):
    """Download a Unicode data file if it is not cached."""

    destination = Path(destination)

    if destination.exists():
        return

    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    print(f"Downloading {url}")

    response = requests.get(
        url,
        timeout=120,
        headers={
            "User-Agent":
                "Unicode String-to-Glyph Assignment Tool"
        },
    )

    response.raise_for_status()

    destination.write_bytes(response.content)


def parse_codepoint_range(value):
    """
    Parse:

        0041

    or:

        0041..005A

    into (start, end).
    """

    value = value.strip()

    if ".." in value:
        start, end = value.split("..", 1)

        return (
            int(start, 16),
            int(end, 16),
        )

    codepoint = int(value, 16)

    return codepoint, codepoint


# ============================================================
# UNICODE DATA
# ============================================================

class UnicodeDatabase:
    """
    Loads the Unicode data required by the assignment system.

    We deliberately use the Unicode Character Database rather
    than Python's built-in unicodedata version, so that the
    assignment rules correspond specifically to Unicode 18.0.
    """

    def __init__(self):
        self.assigned = set()
        self.general_category = {}
        self.east_asian_width = {}

        self.load()

    def load(self):

        DATA_DIRECTORY.mkdir(
            parents=True,
            exist_ok=True,
        )

        unicode_data_file = (
            DATA_DIRECTORY / "UnicodeData.txt"
        )

        east_asian_width_file = (
            DATA_DIRECTORY / "EastAsianWidth.txt"
        )

        download_file(
            UNICODE_DATA_URL,
            unicode_data_file,
        )

        download_file(
            EAST_ASIAN_WIDTH_URL,
            east_asian_width_file,
        )

        self.parse_unicode_data(
            unicode_data_file
        )

        self.parse_east_asian_width(
            east_asian_width_file
        )

    # --------------------------------------------------------
    # UnicodeData.txt
    # --------------------------------------------------------

    def parse_unicode_data(self, filename):

        range_start = None
        range_category = None

        with open(
            filename,
            "r",
            encoding="utf-8",
        ) as file:

            for raw_line in file:

                line = raw_line.rstrip("\n")

                if not line or line.startswith("#"):
                    continue

                fields = line.split(";")

                if len(fields) < 3:
                    continue

                codepoint = int(
                    fields[0],
                    16,
                )

                name = fields[1]
                category = fields[2]

                # UnicodeData.txt represents certain large
                # ranges with:
                #
                # <Something, First>
                # <Something, Last>
                #
                # Unicode specifies this convention in UAX #44.

                if name.endswith(", First>"):

                    range_start = codepoint
                    range_category = category

                    continue

                if name.endswith(", Last>"):

                    if range_start is not None:

                        for cp in range(
                            range_start,
                            codepoint + 1,
                        ):
                            self.assigned.add(cp)
                            self.general_category[cp] = (
                                range_category
                            )

                    range_start = None
                    range_category = None

                    continue

                self.assigned.add(codepoint)
                self.general_category[codepoint] = (
                    category
                )

    # --------------------------------------------------------
    # EastAsianWidth.txt
    # --------------------------------------------------------

    def parse_east_asian_width(self, filename):

        with open(
            filename,
            "r",
            encoding="utf-8",
        ) as file:

            for raw_line in file:

                line = raw_line.split(
                    "#",
                    1
                )[0].strip()

                if not line:
                    continue

                if ";" not in line:
                    continue

                codepoints, width = [
                    x.strip()
                    for x in line.split(
                        ";",
                        1
                    )
                ]

                start, end = parse_codepoint_range(
                    codepoints
                )

                for cp in range(
                    start,
                    end + 1,
                ):
                    self.east_asian_width[cp] = width

    # --------------------------------------------------------
    # Eligibility
    # --------------------------------------------------------

    def is_eligible(self, codepoint):
        """
        Determine whether a code point may be automatically
        selected as a glyph.

        Allowed East Asian Width:
            N  = Neutral
            Na = Narrow
            H  = Halfwidth

        Excluded:
            A  = Ambiguous
            W  = Wide
            F  = Fullwidth

        Also excluded:
            marks
            controls
            formatting characters
            surrogates
            private-use characters
            unassigned characters
        """

        if codepoint not in self.assigned:
            return False

        if not 0 <= codepoint <= 0x10FFFF:
            return False

        # UTF-16 surrogate range
        if 0xD800 <= codepoint <= 0xDFFF:
            return False

        category = self.general_category.get(
            codepoint
        )

        if category is None:
            return False

        # M = all combining marks:
        #
        # Mn Nonspacing Mark
        # Mc Spacing Combining Mark
        # Me Enclosing Mark
        #
        # This excludes ordinary combining diacritics and
        # similar characters.
        if category.startswith("M"):
            return False

        # Controls, formatting characters, surrogates,
        # private-use, and unassigned.
        if category in {
            "Cc",
            "Cf",
            "Cs",
            "Co",
            "Cn",
        }:
            return False

        width = self.east_asian_width.get(
            codepoint
        )

        if width not in {
            "N",
            "Na",
            "H",
        }:
            return False

        return True

    # --------------------------------------------------------
    # First available glyph
    # --------------------------------------------------------

    def first_available(self, retired):
        """
        Return the first eligible Unicode character not
        currently retired.
        """

        for codepoint in sorted(
            self.assigned
        ):

            if codepoint in retired:
                continue

            if self.is_eligible(
                codepoint
            ):
                return chr(codepoint)

        raise RuntimeError(
            "No eligible Unicode glyphs remain."
        )


# ============================================================
# WORKBOOK
# ============================================================

def create_assignment_workbook(filename):
    """Create an empty four-column assignment workbook."""

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Assignments"

    sheet.append(HEADERS)

    workbook.save(filename)

    return workbook


def open_assignment_workbook(filename):
    """
    Open an existing assignment workbook or create one.
    """

    filename = Path(filename)

    if filename.exists():
        workbook = load_workbook(filename)

    else:
        workbook = create_assignment_workbook(
            filename
        )

    sheet = workbook.active

    normalize_headers(sheet)

    return workbook


def normalize_headers(sheet):
    """
    Ensure the worksheet has exactly the desired headers.

    Existing assignment workbooks must have GLYPH/NO./STRING/DEF.
    """

    existing = [
        sheet.cell(
            row=1,
            column=i,
        ).value
        for i in range(
            1,
            5,
        )
    ]

    if existing != HEADERS:

        raise ValueError(
            "Assignment workbook must have the headers:\n"
            "GLYPH | NO. | STRING | DEF"
        )


# ============================================================
# ASSIGNMENT LOOKUPS
# ============================================================

def get_rows(sheet):
    """
    Yield worksheet rows as:

        row_number, glyph, number, string, definition
    """

    for row_number in range(
        2,
        sheet.max_row + 1,
    ):

        glyph = sheet.cell(
            row=row_number,
            column=1,
        ).value

        number = sheet.cell(
            row=row_number,
            column=2,
        ).value

        string = sheet.cell(
            row=row_number,
            column=3,
        ).value

        definition = sheet.cell(
            row=row_number,
            column=4,
        ).value

        yield (
            row_number,
            glyph,
            number,
            string,
            definition,
        )


def glyph_row(sheet, glyph):
    """
    Return the row number assigned to glyph, or None.
    """

    for (
        row_number,
        existing_glyph,
        _number,
        _string,
        _definition,
    ) in get_rows(sheet):

        if existing_glyph == glyph:
            return row_number

    return None


def string_rows(sheet, string):
    """Return all rows containing a given STRING."""

    rows = []

    for (
        row_number,
        _glyph,
        _number,
        existing_string,
        _definition,
    ) in get_rows(sheet):

        if existing_string == string:
            rows.append(row_number)

    return rows


def retired_codepoints(sheet):
    """
    Return all code points currently represented in GLYPH.
    """

    retired = set()

    for (
        _row,
        glyph,
        _number,
        _string,
        _definition,
    ) in get_rows(sheet):

        if not glyph:
            continue

        glyph = str(glyph)

        # GLYPH must represent exactly one Unicode
        # code point.
        if len(glyph) != 1:
            raise ValueError(
                f"Invalid GLYPH {glyph!r}: "
                "GLYPH must contain exactly one "
                "Unicode code point."
            )

        retired.add(ord(glyph))

    return retired


# ============================================================
# ROW OPERATIONS
# ============================================================

def set_row(
    sheet,
    row_number,
    glyph,
    string,
    definition,
):
    """Write a complete assignment row."""

    codepoint = ord(glyph)

    sheet.cell(
        row=row_number,
        column=1,
    ).value = glyph

    sheet.cell(
        row=row_number,
        column=2,
    ).value = f"U+{codepoint:04X}"

    sheet.cell(
        row=row_number,
        column=3,
    ).value = string

    sheet.cell(
        row=row_number,
        column=4,
    ).value = definition


def append_row(
    sheet,
    glyph,
    string,
    definition,
):
    """Append a new assignment row."""

    row_number = sheet.max_row + 1

    set_row(
        sheet,
        row_number,
        glyph,
        string,
        definition,
    )

    return row_number


# ============================================================
# ASSIGNMENT ENGINE
# ============================================================

class AssignmentManager:

    def __init__(
        self,
        workbook_path,
        unicode_db,
    ):

        self.workbook_path = Path(
            workbook_path
        )

        self.unicode = unicode_db

        self.workbook = open_assignment_workbook(
            self.workbook_path
        )

        self.sheet = self.workbook.active

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    def save(self):
        self.workbook.save(
            self.workbook_path
        )

    # --------------------------------------------------------
    # Automatic string assignment
    # --------------------------------------------------------

    def assign_string(
        self,
        string,
        definition="",
    ):
        """
        Assign a fresh glyph to a multi-character string.
        """

        if not string:
            raise ValueError(
                "STRING cannot be empty."
            )

        if len(string) == 1:
            return self.assign_explicit_character(
                string,
                definition,
            )

        # Prevent duplicate semantic assignments.
        existing = string_rows(
            self.sheet,
            string,
        )

        if existing:
            raise ValueError(
                f"STRING {string!r} is already "
                f"assigned on row {existing[0]}."
            )

        retired = retired_codepoints(
            self.sheet
        )

        glyph = self.unicode.first_available(
            retired
        )

        append_row(
            self.sheet,
            glyph,
            string,
            definition,
        )

        self.save()

        return glyph

    # --------------------------------------------------------
    # Explicit character assignment
    # --------------------------------------------------------

    def assign_explicit_character(
        self,
        character,
        definition="",
    ):
        """
        Make a single Unicode character own itself.

        If the glyph is already assigned, move the existing
        assignment to a fresh glyph.

        The NEW explicit assignment receives STRING =
        character and the supplied DEF.
        """

        if len(character) != 1:
            raise ValueError(
                "Explicit assignment requires exactly "
                "one Unicode code point."
            )

        codepoint = ord(character)

        # ----------------------------------------------------
        # Locate existing owner.
        # ----------------------------------------------------

        existing_row = glyph_row(
            self.sheet,
            character,
        )

        # ----------------------------------------------------
        # Character is unused.
        # ----------------------------------------------------

        if existing_row is None:

            append_row(
                self.sheet,
                character,
                character,
                definition,
            )

            self.save()

            return character

        # ----------------------------------------------------
        # Character is already owned.
        #
        # Preserve the old STRING and DEF.
        # ----------------------------------------------------

        old_string = self.sheet.cell(
            row=existing_row,
            column=3,
        ).value

        old_definition = self.sheet.cell(
            row=existing_row,
            column=4,
        ).value

        # Current assignments are retired, except the glyph
        # we are about to take.
        retired = retired_codepoints(
            self.sheet
        )

        retired.remove(codepoint)

        # Find a replacement BEFORE changing the existing row.
        replacement = self.unicode.first_available(
            retired
        )

        # Move the old assignment.
        set_row(
            self.sheet,
            existing_row,
            replacement,
            old_string,
            old_definition,
        )

        # Add the explicit assignment.
        append_row(
            self.sheet,
            character,
            character,
            definition,
        )

        self.save()

        return character

    # --------------------------------------------------------
    # Generic input
    # --------------------------------------------------------

    def assign(
        self,
        value,
        definition="",
    ):
        """
        Automatically choose behavior based on the length
        of the input.

        One Unicode code point:
            claim that exact glyph.

        More than one:
            allocate a new glyph.
        """

        if not value:
            raise ValueError(
                "Input cannot be empty."
            )

        if len(value) == 1:
            return self.assign_explicit_character(
                value,
                definition,
            )

        return self.assign_string(
            value,
            definition,
        )

    # --------------------------------------------------------
    # Show assignment
    # --------------------------------------------------------

    def show(self, string):
        """Print all rows associated with STRING."""

        found = False

        for (
            row_number,
            glyph,
            number,
            existing_string,
            definition,
        ) in get_rows(self.sheet):

            if existing_string == string:

                found = True

                print()
                print(
                    f"Row:    {row_number}"
                )
                print(
                    f"Glyph:  {glyph}"
                )
                print(
                    f"NO.:    {number}"
                )
                print(
                    f"STRING: {existing_string}"
                )
                print(
                    f"DEF:    {definition or ''}"
                )
                print()

        if not found:
            print(
                f"No assignment found for {string!r}."
            )

    # --------------------------------------------------------
    # Count
    # --------------------------------------------------------

    def count(self):

        retired = retired_codepoints(
            self.sheet
        )

        available = 0

        for codepoint in self.unicode.assigned:

            if codepoint in retired:
                continue

            if self.unicode.is_eligible(
                codepoint
            ):
                available += 1

        print(
            f"Assigned glyphs: "
            f"{len(retired):,}"
        )

        print(
            f"Available eligible glyphs: "
            f"{available:,}"
        )

        print(
            f"Rows: "
            f"{max(0, self.sheet.max_row - 1):,}"
        )


# ============================================================
# SPREADSHEET IMPORT
# ============================================================

def locate_column(headers, name):
    """
    Find a column by header name, case-insensitively.
    """

    target = name.strip().upper()

    for index, header in enumerate(
        headers,
        start=1,
    ):

        if header is None:
            continue

        if str(header).strip().upper() == target:
            return index

    return None


def import_spreadsheet(
    source_path,
    output_path,
    unicode_db,
):
    """
    Import a prepopulated spreadsheet containing STRING and DEF.

    The source may have:

        STRING | DEF

    or may already have:

        GLYPH | NO. | STRING | DEF

    For a fresh import, STRING/DEF are treated as the source
    data and glyphs are assigned in row order.
    """

    source_path = Path(
        source_path
    )

    output_path = Path(
        output_path
    )

    if not source_path.exists():
        raise FileNotFoundError(
            f"Input spreadsheet not found: "
            f"{source_path}"
        )

    source_workbook = load_workbook(
        source_path
    )

    source_sheet = (
        source_workbook.active
    )

    headers = [
        source_sheet.cell(
            row=1,
            column=column,
        ).value
        for column in range(
            1,
            source_sheet.max_column + 1,
        )
    ]

    string_column = locate_column(
        headers,
        "STRING",
    )

    def_column = locate_column(
        headers,
        "DEF",
    )

    if string_column is None:
        raise ValueError(
            "Input spreadsheet must contain "
            "a STRING column."
        )

    if def_column is None:
        raise ValueError(
            "Input spreadsheet must contain "
            "a DEF column."
        )

    # --------------------------------------------------------
    # Create a fresh output workbook.
    # --------------------------------------------------------

    if output_path.exists():

        backup_path = output_path.with_suffix(
            output_path.suffix + ".bak"
        )

        shutil.copy2(
            output_path,
            backup_path,
        )

        print(
            f"Existing output backed up to "
            f"{backup_path}"
        )

    create_assignment_workbook(
        output_path
    )

    manager = AssignmentManager(
        output_path,
        unicode_db,
    )

    imported = 0

    for row_number in range(
        2,
        source_sheet.max_row + 1,
    ):

        string_value = source_sheet.cell(
            row=row_number,
            column=string_column,
        ).value

        definition = source_sheet.cell(
            row=row_number,
            column=def_column,
        ).value

        if string_value is None:
            continue

        string_value = str(
            string_value
        )

        if not string_value:
            continue

        if definition is None:
            definition = ""

        definition = str(
            definition
        )

        try:

            glyph = manager.assign(
                string_value,
                definition,
            )

            imported += 1

            print(
                f"[{imported}] "
                f"{glyph} -> "
                f"{string_value}"
            )

        except Exception as error:

            raise RuntimeError(
                f"Could not import row "
                f"{row_number}: "
                f"{error}"
            ) from error

    manager.save()

    print()
    print(
        f"Import complete: "
        f"{imported:,} rows."
    )

    print(
        f"Output: {output_path}"
    )


# ============================================================
# INTERACTIVE SHELL
# ============================================================

def interactive(manager):

    print()
    print(
        "Unicode String-to-Glyph Assignment Tool"
    )
    print(
        f"Unicode version: {UNICODE_VERSION}"
    )
    print(
        f"Workbook: {manager.workbook_path}"
    )
    print()
    print(
        "Enter a string to assign a glyph."
    )
    print(
        "Enter one character to claim that exact glyph."
    )
    print()
    print(
        "Commands:"
    )
    print(
        "  :import FILE       Import STRING/DEF spreadsheet"
    )
    print(
        "  :show STRING       Show an assignment"
    )
    print(
        "  :count             Show assignment statistics"
    )
    print(
        "  :save              Save workbook"
    )
    print(
        "  :quit              Exit"
    )
    print()

    while True:

        try:
            value = input("> ")

        except EOFError:
            print()
            break

        except KeyboardInterrupt:
            print()
            break

        if not value:
            continue

        # ----------------------------------------------------
        # Quit
        # ----------------------------------------------------

        if value == ":quit":
            manager.save()
            break

        # ----------------------------------------------------
        # Save
        # ----------------------------------------------------

        if value == ":save":
            manager.save()
            print("Saved.")
            continue

        # ----------------------------------------------------
        # Count
        # ----------------------------------------------------

        if value == ":count":
            manager.count()
            continue

        # ----------------------------------------------------
        # Show
        # ----------------------------------------------------

        if value.startswith(":show "):

            string = value[
                len(":show "):
            ]

            manager.show(string)

            continue

        # ----------------------------------------------------
        # Import
        # ----------------------------------------------------

        if value.startswith(":import "):

            filename = value[
                len(":import "):
            ].strip()

            if not filename:
                print(
                    "Usage: :import filename.xlsx"
                )
                continue

            try:

                import_spreadsheet(
                    filename,
                    manager.workbook_path,
                    manager.unicode,
                )

                # Reload after import.
                manager.workbook = (
                    load_workbook(
                        manager.workbook_path
                    )
                )

                manager.sheet = (
                    manager.workbook.active
                )

            except Exception as error:

                print(
                    f"IMPORT ERROR: {error}"
                )

            continue

        # ----------------------------------------------------
        # Normal assignment
        # ----------------------------------------------------

        try:

            glyph = manager.assign(value)

            print(
                f"Assigned {glyph} "
                f"(U+{ord(glyph):04X}) "
                f"-> {value!r}"
            )

        except Exception as error:

            print(
                f"ERROR: {error}"
            )


# ============================================================
# COMMAND-LINE INTERFACE
# ============================================================

def main():

    parser = argparse.ArgumentParser(
        description=(
            "Assign Unicode glyphs to strings and "
            "maintain the assignments in Excel."
        )
    )

    parser.add_argument(
        "--import",
        dest="import_file",
        metavar="FILE",
        help=(
            "Import a spreadsheet containing "
            "STRING and DEF columns."
        ),
    )

    parser.add_argument(
        "--output",
        default=str(DEFAULT_WORKBOOK),
        metavar="FILE",
        help=(
            "Assignment workbook to create/use. "
            f"Default: {DEFAULT_WORKBOOK}"
        ),
    )

    args = parser.parse_args()

    try:

        print(
            f"Loading Unicode {UNICODE_VERSION} data..."
        )

        unicode_db = UnicodeDatabase()

        print(
            f"Loaded "
            f"{len(unicode_db.assigned):,} "
            f"assigned Unicode code points."
        )

        print()

        if args.import_file:

            import_spreadsheet(
                args.import_file,
                args.output,
                unicode_db,
            )

            print()

        manager = AssignmentManager(
            args.output,
            unicode_db,
        )

        interactive(manager)

    except KeyboardInterrupt:

        print()
        print("Stopped.")

    except Exception as error:

        print(
            f"FATAL ERROR: {error}",
            file=sys.stderr,
        )

        sys.exit(1)


if __name__ == "__main__":
    main()
```
