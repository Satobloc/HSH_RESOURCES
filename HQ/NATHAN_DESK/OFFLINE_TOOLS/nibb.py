from pathlib import Path
import json
import re

# ============================================================
# CONFIG
# ============================================================

SOURCE_DIR = Path(r"C:\Users\Admin\Dropbox\PC\Downloads\BITE_SIZED")
OUTPUT_DIR = SOURCE_DIR / "_TXT_CHUNKS"

# Stay slightly under 2.5 MB for safety.
MAX_BYTES = 2_450_000

# File types that can normally just be read as text.
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown",
    ".json", ".jsonl", ".ndjson",
    ".csv", ".tsv",
    ".xml", ".html", ".htm",
    ".yaml", ".yml",
    ".log",
    ".rtf",
    ".py", ".js", ".ts",
    ".css",
    ".tex",
    ".ini", ".cfg", ".conf",
}

SKIP_DIRS = {
    "_TXT_CHUNKS",
    ".git",
    "__pycache__",
}


# ============================================================
# TEXT EXTRACTION
# ============================================================

def read_plain_text(path):
    """
    Try several common encodings rather than dying on one odd file.
    """
    encodings = [
        "utf-8",
        "utf-8-sig",
        "utf-16",
        "cp1252",
        "latin-1",
    ]

    for encoding in encodings:
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
        except Exception as exc:
            print(f"  ERROR reading {path.name}: {exc}")
            return None

    return None


def extract_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        print("  PDF skipped: install pypdf with: pip install pypdf")
        return None

    try:
        reader = PdfReader(path)

        pieces = []

        for page_num, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                text = f"\n[PDF PAGE {page_num} EXTRACTION ERROR: {exc}]\n"

            pieces.append(
                f"\n\n===== PAGE {page_num} =====\n\n{text}"
            )

        return "".join(pieces)

    except Exception as exc:
        print(f"  PDF ERROR: {exc}")
        return None


def extract_docx(path):
    try:
        from docx import Document
    except ImportError:
        print("  DOCX skipped: install python-docx with: pip install python-docx")
        return None

    try:
        doc = Document(path)

        parts = []

        for paragraph in doc.paragraphs:
            parts.append(paragraph.text)

        # Also get text from tables.
        for table_num, table in enumerate(doc.tables, start=1):
            parts.append(f"\n===== TABLE {table_num} =====")

            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                parts.append("\t".join(cells))

        return "\n".join(parts)

    except Exception as exc:
        print(f"  DOCX ERROR: {exc}")
        return None


def extract_json(path):
    """
    JSON is normalized into pretty-printed text if possible.
    If malformed or unusual, falls back to reading it verbatim.
    """
    raw = read_plain_text(path)

    if raw is None:
        return None

    try:
        obj = json.loads(raw)

        return json.dumps(
            obj,
            ensure_ascii=False,
            indent=2
        )

    except Exception:
        return raw


def extract_text(path):
    ext = path.suffix.lower()

    if ext == ".pdf":
        return extract_pdf(path)

    if ext == ".docx":
        return extract_docx(path)

    if ext == ".json":
        return extract_json(path)

    if ext in TEXT_EXTENSIONS:
        return read_plain_text(path)

    # Last-resort attempt for unfamiliar file extensions.
    # This lets us catch text-containing files that simply have
    # an extension we didn't anticipate.
    try:
        data = path.read_bytes()

        # Binary files usually contain NUL bytes.
        if b"\x00" in data[:10000]:
            return None

        for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue

    except Exception:
        pass

    return None


# ============================================================
# UTF-8-SAFE CHUNKING
# ============================================================

def split_long_string_by_bytes(text, byte_limit):
    """
    Split a single enormous string safely without cutting a UTF-8
    character in half.
    """
    pieces = []
    remaining = text

    while len(remaining.encode("utf-8")) > byte_limit:

        # Binary search for largest character count fitting byte limit.
        low = 1
        high = len(remaining)

        while low <= high:
            mid = (low + high) // 2
            size = len(remaining[:mid].encode("utf-8"))

            if size <= byte_limit:
                low = mid + 1
            else:
                high = mid - 1

        cut = max(1, high)

        # Prefer a sensible break point if one exists reasonably nearby.
        candidate = remaining[:cut]

        break_positions = [
            candidate.rfind("\n\n"),
            candidate.rfind("\n"),
            candidate.rfind(". "),
            candidate.rfind(" "),
        ]

        best_break = max(break_positions)

        # Don't sacrifice a huge portion just to find a delimiter.
        if best_break > cut * 0.75:
            cut = best_break + 1

        pieces.append(remaining[:cut])
        remaining = remaining[cut:]

    if remaining:
        pieces.append(remaining)

    return pieces


def chunk_text(text, max_bytes):
    """
    Prefer splitting along existing lines while guaranteeing each
    resulting string stays under max_bytes in UTF-8.
    """
    chunks = []
    current_parts = []
    current_size = 0

    lines = text.splitlines(keepends=True)

    for line in lines:

        encoded_size = len(line.encode("utf-8"))

        # A single line may itself exceed the limit.
        if encoded_size > max_bytes:

            if current_parts:
                chunks.append("".join(current_parts))
                current_parts = []
                current_size = 0

            chunks.extend(
                split_long_string_by_bytes(line, max_bytes)
            )

            continue

        if current_size + encoded_size > max_bytes:

            if current_parts:
                chunks.append("".join(current_parts))

            current_parts = [line]
            current_size = encoded_size

        else:
            current_parts.append(line)
            current_size += encoded_size

    if current_parts:
        chunks.append("".join(current_parts))

    # Edge case: text with no line endings / empty splitlines behavior.
    if not chunks and text:
        chunks = split_long_string_by_bytes(text, max_bytes)

    return chunks


# ============================================================
# OUTPUT
# ============================================================

def sanitize_filename(name):
    return re.sub(r'[<>:"/\\|?*]', "_", name)


def process_file(path):
    relative = path.relative_to(SOURCE_DIR)

    print(f"\nREADING: {relative}")

    text = extract_text(path)

    if text is None:
        print("  SKIPPED — unsupported/binary/unreadable")
        return 0, 0

    if not text.strip():
        print("  SKIPPED — no extractable text")
        return 0, 0

    chunks = chunk_text(text, MAX_BYTES)

    # Mirror original subdirectory structure.
    relative_parent = relative.parent
    destination_dir = OUTPUT_DIR / relative_parent
    destination_dir.mkdir(parents=True, exist_ok=True)

    stem = sanitize_filename(path.stem)

    total = len(chunks)

    for index, chunk in enumerate(chunks, start=1):

        filename = (
            f"{stem}__chunk_{index:04d}_of_{total:04d}.txt"
        )

        output_path = destination_dir / filename

        output_path.write_text(
            chunk,
            encoding="utf-8",
            newline=""
        )

        size = output_path.stat().st_size

        print(
            f"  -> {filename} "
            f"({size / 1_000_000:.3f} MB)"
        )

        if size >= 2_500_000:
            print("     WARNING: chunk exceeded requested ceiling!")

    return 1, total


# ============================================================
# MAIN
# ============================================================

def main():

    if not SOURCE_DIR.exists():
        print(f"Source folder does not exist:\n{SOURCE_DIR}")
        return

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    files_seen = 0
    files_processed = 0
    chunks_written = 0

    for path in SOURCE_DIR.rglob("*"):

        if not path.is_file():
            continue

        if any(part in SKIP_DIRS for part in path.parts):
            continue

        files_seen += 1

        processed, chunks = process_file(path)

        files_processed += processed
        chunks_written += chunks

    print("\n" + "=" * 60)
    print("DONE")
    print("=" * 60)
    print(f"Files examined:  {files_seen}")
    print(f"Files converted: {files_processed}")
    print(f"TXT chunks:      {chunks_written}")
    print(f"Output folder:   {OUTPUT_DIR}")


if __name__ == "__main__":
    main()